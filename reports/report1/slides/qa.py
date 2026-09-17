"""Kiểm tra hình học và nội dung deck — thay cho QA bằng mắt.

Máy này không có LibreOffice và COM automation của PowerPoint lỗi type library,
nên không render được ảnh từng slide. Script này kiểm những lỗi mà QA bằng mắt
thường bắt: tràn lề, chữ vượt khung, các khối đè nhau, chữ còn sót placeholder.

    python qa.py DSP391m_Report1_Proposal.pptx
"""
from __future__ import annotations

import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
SAFE_MARGIN = 0.5      # inch — lề tối thiểu cho nội dung
# Hệ số ước lượng: bề rộng trung bình một ký tự so với cỡ chữ, cho Calibri/Cambria
CHAR_W = 0.50
LINE_H = 1.22          # chiều cao dòng so với cỡ chữ


def inches(v) -> float:
    return Emu(int(v)).inches if v is not None else 0.0


def est_text_height(text: str, font_pt: float, box_w_in: float) -> float:
    """Ước lượng chiều cao khối chữ, tính cả ngắt dòng."""
    if not text.strip():
        return 0.0
    char_w_in = font_pt * CHAR_W / 72
    per_line = max(int(box_w_in / char_w_in), 1)
    lines = 0
    for para in text.split("\n"):
        lines += max(1, -(-len(para) // per_line))
    return lines * font_pt * LINE_H / 72


def main(path: str) -> int:
    prs = Presentation(path)
    SW, SH = prs.slide_width.inches, prs.slide_height.inches
    print(f"{path}\nKhổ slide: {SW:.2f} × {SH:.2f} in · {len(prs.slides)} slide\n")

    issues: list[str] = []
    placeholders = ("lorem", "ipsum", "todo", "[insert", "xxx")

    for i, slide in enumerate(prs.slides, 1):
        boxes = []
        for sh in slide.shapes:
            x, y = inches(sh.left), inches(sh.top)
            w, h = inches(sh.width), inches(sh.height)

            # 1. vượt ra ngoài slide
            if x < -0.02 or y < -0.02 or x + w > SW + 0.02 or y + h > SH + 0.02:
                issues.append(f"  slide {i}: '{sh.shape_type}' ra ngoài khổ "
                              f"({x:.2f},{y:.2f} {w:.2f}×{h:.2f})")

            if not sh.has_text_frame:
                continue
            text = sh.text_frame.text
            if not text.strip():
                continue

            # 2. placeholder còn sót
            low = text.lower()
            for p in placeholders:
                if p in low:
                    issues.append(f"  slide {i}: còn chữ placeholder '{p}'")

            # 3. chữ vượt khung
            sizes = [r.font.size.pt for para in sh.text_frame.paragraphs
                     for r in para.runs if r.font.size]
            if sizes and w > 0.2:
                need = est_text_height(text, max(sizes), w)
                if need > h * 1.55:      # nới rộng vì ước lượng thô
                    issues.append(
                        f"  slide {i}: nghi tràn chữ — cần ~{need:.2f} in, "
                        f"khung {h:.2f} in · \"{text[:46]}…\"")

            boxes.append((x, y, w, h, text[:28]))

        # 4. hai khối chữ đè nhau
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                ax, ay, aw, ah, at = boxes[a]
                bx, by, bw, bh, bt = boxes[b]
                ox = min(ax + aw, bx + bw) - max(ax, bx)
                oy = min(ay + ah, by + bh) - max(ay, by)
                if ox > 0.12 and oy > 0.12:
                    issues.append(f"  slide {i}: chữ đè nhau {ox:.2f}×{oy:.2f} in "
                                  f"— \"{at}\" ↔ \"{bt}\"")

    if issues:
        print(f"⚠️  {len(issues)} điểm cần xem:")
        for s in issues:
            print(s)
    else:
        print("✅ Không thấy tràn lề, tràn chữ, đè nhau hay placeholder sót.")

    print("\n--- Nội dung từng slide ---")
    for i, slide in enumerate(prs.slides, 1):
        heads = [sh.text_frame.text.split("\n")[0]
                 for sh in slide.shapes
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        n_notes = (len(slide.notes_slide.notes_text_frame.text.split())
                   if slide.has_notes_slide else 0)
        print(f"{i:2}. {heads[1] if len(heads) > 1 else heads[0]:<52} "
              f"({len(slide.shapes)} khối, ghi chú {n_notes} từ)")
    return 1 if issues else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "DSP391m_Report1_Proposal.pptx"))
